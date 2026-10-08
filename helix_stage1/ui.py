from __future__ import annotations

import logging
import queue
import re
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Literal
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from .automation import AutomationWorker, WorkerCommand, WorkerEvent
from .config import LOG_DIR
from .resolver_catalog import RESOLVER_USERS


INCIDENT_PATTERN = re.compile(r"INC\d{12}", re.IGNORECASE)
PARTIAL_INCIDENT_PATTERN = re.compile(r"(?:I(?:N(?:C\d{0,12})?)?)?", re.IGNORECASE)
INCIDENT_HEADER_NAMES = {"incidencia", "incidencias"}
REPORT_HEADER = "numero de incidencia      Estado"
REPORT_TITLE_FORMAT = "%m/%d/%Y   %H:%M hrs"
BatchStatus = Literal["OK", "ERROR"]


def _create_logger() -> logging.Logger:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("helix_stage1")
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        handler = logging.FileHandler(LOG_DIR / "automation.log", encoding="utf-8")
        handler.setFormatter(
            logging.Formatter("%(asctime)s %(levelname)s %(threadName)s %(message)s")
        )
        logger.addHandler(handler)
    return logger


def _report_now() -> datetime:
    try:
        return datetime.now(ZoneInfo("America/Santiago"))
    except ZoneInfoNotFoundError:
        return datetime.now()


class HelixApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.commands: queue.Queue[WorkerCommand] = queue.Queue()
        self.events: queue.Queue[WorkerEvent] = queue.Queue()
        self.worker = AutomationWorker(self.commands, self.events, _create_logger())
        self.closing = False
        self.busy = True
        self.loaded_incidents: list[str] = []
        self.batch_running = False
        self.batch_index = 0
        self.batch_results: list[tuple[str, BatchStatus, str]] = []

        root.title("Consulta de incidencias | BMC Helix")
        root.geometry("760x680")
        root.minsize(680, 560)
        root.protocol("WM_DELETE_WINDOW", self._on_close)

        style = ttk.Style(root)
        style.configure("Title.TLabel", font=("Segoe UI", 11, "bold"))
        style.configure("Section.TLabelframe.Label", font=("Segoe UI", 10, "bold"))
        style.configure("Status.TLabel", font=("Segoe UI", 9, "bold"))

        container = ttk.Frame(root, padding=12)
        container.pack(fill="both", expand=True)
        container.columnconfigure(0, weight=1)
        container.rowconfigure(5, weight=1)

        ttk.Label(container, text="Cierre de incidencias Helix", style="Title.TLabel").grid(
            row=0, column=0, sticky="w"
        )

        assignment_frame = ttk.LabelFrame(
            container, text="1. Datos comunes de cierre", style="Section.TLabelframe"
        )
        assignment_frame.grid(row=1, column=0, sticky="ew", pady=(14, 10))
        assignment_frame.columnconfigure(0, weight=1, uniform="common")
        assignment_frame.columnconfigure(1, weight=1, uniform="common")

        ttk.Label(assignment_frame, text="Grupo Resolutor").grid(
            row=0, column=0, sticky="w", pady=(0, 4), padx=(8, 6)
        )
        ttk.Label(assignment_frame, text="Usuario Resolutor").grid(
            row=0, column=1, sticky="w", pady=(0, 4), padx=(6, 8)
        )
        self.resolver_group_var = tk.StringVar()
        self.resolver_group_combo = ttk.Combobox(
            assignment_frame,
            textvariable=self.resolver_group_var,
            values=tuple(RESOLVER_USERS),
            state="readonly",
        )
        self.resolver_group_combo.grid(row=1, column=0, sticky="ew", padx=(8, 6))
        self.resolver_group_combo.bind("<<ComboboxSelected>>", self._on_group_selected)

        self.resolver_user_var = tk.StringVar()
        self.resolver_user_combo = ttk.Combobox(
            assignment_frame,
            textvariable=self.resolver_user_var,
            state="disabled",
        )
        self.resolver_user_combo.grid(row=1, column=1, sticky="ew", padx=(6, 8))

        ttk.Label(assignment_frame, text="Resolución").grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(8, 4), padx=8
        )
        self.resolution_text = tk.Text(assignment_frame, wrap="word", height=3)
        self.resolution_text.grid(row=3, column=0, columnspan=2, sticky="ew", padx=8, pady=(0, 8))

        modes_frame = ttk.Frame(container)
        modes_frame.grid(row=2, column=0, sticky="ew")
        modes_frame.columnconfigure(0, weight=1, uniform="mode")
        modes_frame.columnconfigure(1, weight=1, uniform="mode")

        manual_frame = ttk.LabelFrame(
            modes_frame, text="2A. Cierre manual", style="Section.TLabelframe"
        )
        manual_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        manual_frame.columnconfigure(0, weight=1)

        ttk.Label(manual_frame, text="Número de incidencia").grid(
            row=0, column=0, sticky="w", pady=(0, 4), padx=8
        )
        self.incident_var = tk.StringVar()
        validator = root.register(self._validate_partial_incident)
        self.entry = ttk.Entry(
            manual_frame,
            textvariable=self.incident_var,
            validate="key",
            validatecommand=(validator, "%P"),
            width=32,
        )
        self.entry.grid(row=1, column=0, sticky="ew", padx=8)
        self.entry.bind("<Return>", self._on_enter)

        self.search_button = ttk.Button(
            manual_frame,
            text="Cierre manual",
            command=self._submit,
            state="disabled",
        )
        self.search_button.grid(row=2, column=0, sticky="ew", padx=8, pady=(8, 8))

        excel_frame = ttk.LabelFrame(
            modes_frame, text="2B. Cierre masivo por Excel", style="Section.TLabelframe"
        )
        excel_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
        excel_frame.columnconfigure(0, weight=1)
        excel_frame.columnconfigure(1, weight=1)

        self.load_excel_button = ttk.Button(
            excel_frame,
            text="Cargar Excel",
            command=self._load_excel,
            state="disabled",
        )
        self.load_excel_button.grid(row=0, column=0, sticky="ew", padx=(8, 4), pady=(8, 6))

        self.process_excel_button = ttk.Button(
            excel_frame,
            text="Cierre masivo",
            command=self._start_excel_batch,
            state="disabled",
        )
        self.process_excel_button.grid(row=0, column=1, sticky="ew", padx=(4, 8), pady=(8, 6))

        self.excel_status_var = tk.StringVar(value="Sin Excel cargado.")
        ttk.Label(excel_frame, textvariable=self.excel_status_var, wraplength=260).grid(
            row=1, column=0, columnspan=2, sticky="ew", padx=8, pady=(0, 8)
        )
        self.batch_progress_var = tk.StringVar(value="Sin proceso activo.")
        self.batch_progress = ttk.Progressbar(
            excel_frame,
            mode="determinate",
            maximum=1,
            value=0,
        )
        self.batch_progress.grid(row=2, column=0, columnspan=2, sticky="ew", padx=8)
        ttk.Label(excel_frame, textvariable=self.batch_progress_var, wraplength=260).grid(
            row=3, column=0, columnspan=2, sticky="ew", padx=8, pady=(4, 8)
        )

        self.status_var = tk.StringVar(value="Iniciando la interfaz…")
        status_frame = ttk.LabelFrame(container, text="3. Estado", style="Section.TLabelframe")
        status_frame.grid(row=3, column=0, sticky="ew", pady=(10, 10))
        status_frame.columnconfigure(0, weight=1)
        ttk.Label(
            status_frame,
            textvariable=self.status_var,
            style="Status.TLabel",
            wraplength=660,
        ).grid(
            row=0, column=0, sticky="ew", padx=8, pady=6
        )

        results_frame = ttk.LabelFrame(container, text="Resultados", style="Section.TLabelframe")
        results_frame.grid(row=4, column=0, sticky="ew", pady=(0, 10))
        results_frame.columnconfigure(0, weight=1)
        self.results_table = ttk.Treeview(
            results_frame,
            columns=("index", "incident", "status", "detail"),
            show="headings",
            height=4,
        )
        self.results_table.heading("index", text="#")
        self.results_table.heading("incident", text="Incidencia")
        self.results_table.heading("status", text="Estado")
        self.results_table.heading("detail", text="Detalle")
        self.results_table.column("index", width=42, anchor="center", stretch=False)
        self.results_table.column("incident", width=150, stretch=False)
        self.results_table.column("status", width=82, anchor="center", stretch=False)
        self.results_table.column("detail", width=420, stretch=True)
        self.results_table.grid(row=0, column=0, sticky="ew", padx=8, pady=6)

        log_frame = ttk.LabelFrame(container, text="Registro", style="Section.TLabelframe")
        log_frame.grid(row=5, column=0, sticky="nsew")
        log_frame.rowconfigure(0, weight=1)
        log_frame.columnconfigure(0, weight=1)
        self.log_text = tk.Text(log_frame, wrap="word", state="disabled", height=5)
        scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        self.worker.start()
        self.root.after(100, self._poll_events)
        self.entry.focus_set()

    @staticmethod
    def _validate_partial_incident(candidate: str) -> bool:
        return PARTIAL_INCIDENT_PATTERN.fullmatch(candidate) is not None

    def _on_enter(self, _event: tk.Event[tk.Misc]) -> str:
        self._submit()
        return "break"

    def _on_group_selected(self, _event: tk.Event[tk.Misc]) -> None:
        group = self.resolver_group_var.get()
        users = RESOLVER_USERS.get(group, ())
        self.resolver_user_var.set("")
        self.resolver_user_combo.configure(values=users)
        self.resolver_user_combo.configure(state="readonly" if users else "disabled")

    @staticmethod
    def _normalize_header(value: object) -> str:
        text = str(value or "").strip().casefold()
        return re.sub(r"[^a-z0-9]+", "", text)

    def _load_excel(self) -> None:
        if self.busy or self.closing:
            return

        path = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=(
                ("Planillas", "*.xlsx *.xlsm *.ods"),
                ("Archivos Excel", "*.xlsx *.xlsm"),
                ("OpenDocument", "*.ods"),
                ("Todos los archivos", "*.*"),
            ),
        )
        if not path:
            return

        try:
            incidents, duplicates, invalid_values, report_rows = self._read_incidents_from_excel(
                Path(path)
            )
        except Exception as error:
            self.loaded_incidents = []
            self.excel_status_var.set("No se pudo cargar el Excel.")
            self.batch_progress.configure(maximum=1, value=0)
            self.batch_progress_var.set("Sin proceso activo.")
            self._set_idle_buttons()
            self._append_log(f"Error cargando Excel: {error}")
            return

        self.loaded_incidents = incidents
        report_path = self._write_load_report(Path(path), report_rows)
        self.excel_status_var.set(
            f"{len(incidents)} cargadas. Listo para cierre masivo."
        )
        self.batch_progress.configure(maximum=max(1, len(incidents)), value=0)
        self.batch_progress_var.set(f"Pendientes: {len(incidents)}")
        self._set_idle_buttons()
        self._append_log(
            f"Excel cargado: {len(incidents)} incidencias válidas."
        )
        self._append_log(f"Resumen TXT generado: {report_path}")
        if duplicates:
            self._append_log(f"Duplicadas omitidas: {duplicates}")
        if invalid_values:
            preview = ", ".join(invalid_values[:5])
            suffix = "…" if len(invalid_values) > 5 else ""
            self._append_log(f"Valores inválidos omitidos: {preview}{suffix}")

    def _read_incidents_from_excel(
        self, path: Path
    ) -> tuple[list[str], int, list[str], list[tuple[str, str]]]:
        suffix = path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            rows = self._read_spreadsheet_rows_with_openpyxl(path)
        elif suffix == ".ods":
            rows = self._read_spreadsheet_rows_with_ods(path)
        else:
            raise RuntimeError(
                "Formato no soportado. Usa .xlsx, .xlsm o .ods."
            )

        return self._extract_incidents_from_rows(rows)

    def _read_spreadsheet_rows_with_openpyxl(self, path: Path) -> list[tuple[object, ...]]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError(
                "Falta instalar openpyxl. Ejecuta: pip install -r requirements.txt"
            ) from error

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

    def _read_spreadsheet_rows_with_ods(self, path: Path) -> list[tuple[object, ...]]:
        try:
            from pyexcel_ods3 import get_data
        except ImportError as error:
            raise RuntimeError(
                "Falta instalar pyexcel-ods3. Ejecuta: pip install -r requirements.txt"
            ) from error

        sheets = get_data(str(path))
        if not sheets:
            raise RuntimeError("El archivo .ods no contiene hojas.")
        first_sheet_rows = next(iter(sheets.values()))
        return [tuple(row) for row in first_sheet_rows]

    def _extract_incidents_from_rows(
        self, rows: list[tuple[object, ...]]
    ) -> tuple[list[str], int, list[str], list[tuple[str, str]]]:
        header_index: int | None = None
        incident_column: int | None = None

        for row_index, row in enumerate(rows[:10]):
            for column_index, value in enumerate(row):
                if self._normalize_header(value) in INCIDENT_HEADER_NAMES:
                    header_index = row_index
                    incident_column = column_index
                    break
            if incident_column is not None:
                break

        if header_index is None or incident_column is None:
            raise RuntimeError("No encontré una columna con cabecera 'Incidencias'.")

        incidents: list[str] = []
        seen: set[str] = set()
        duplicates = 0
        invalid_values: list[str] = []
        report_rows: list[tuple[str, str]] = []
        for row in rows[header_index + 1 :]:
            raw_value = row[incident_column] if incident_column < len(row) else None
            if raw_value is None:
                continue
            incident = str(raw_value).strip().upper()
            if not incident:
                continue
            if INCIDENT_PATTERN.fullmatch(incident) is None:
                invalid_values.append(incident)
                report_rows.append((incident, "NOK"))
                continue
            if incident in seen:
                duplicates += 1
                report_rows.append((incident, "NOK"))
                continue
            seen.add(incident)
            incidents.append(incident)
            report_rows.append((incident, "OK"))

        if not incidents:
            raise RuntimeError("El archivo no contiene incidencias válidas.")
        return incidents, duplicates, invalid_values, report_rows

    def _write_load_report(self, excel_path: Path, rows: list[tuple[str, str]]) -> Path:
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = re.sub(r"[^A-Za-z0-9_-]", "_", excel_path.stem)[:40] or "excel"
        report_path = LOG_DIR / f"resumen_carga_{timestamp}_{safe_name}.txt"
        with report_path.open("w", encoding="utf-8") as report:
            report.write(REPORT_HEADER + "\n")
            for incident, status in rows:
                report.write(f"{incident:<25}{status}\n")
        return report_path

    def _submit(self) -> None:
        if self.busy or self.closing:
            return
        incident = self.incident_var.get().strip().upper()
        if INCIDENT_PATTERN.fullmatch(incident) is None:
            self.status_var.set("Formato inválido: usa INC seguido de 12 dígitos.")
            self._append_log("Formato inválido. Ejemplo: INC123456789012")
            self.entry.focus_set()
            return

        resolver_group = self.resolver_group_var.get()
        resolver_user = self.resolver_user_var.get()
        if resolver_group not in RESOLVER_USERS:
            self.status_var.set("Selecciona un Grupo Resolutor.")
            self._append_log("Falta seleccionar el Grupo Resolutor.")
            self.resolver_group_combo.focus_set()
            return
        if resolver_user not in RESOLVER_USERS[resolver_group]:
            self.status_var.set("Selecciona un usuario válido para el grupo elegido.")
            self._append_log("Falta seleccionar un Usuario Resolutor válido para el grupo.")
            self.resolver_user_combo.focus_set()
            return
        resolution = self.resolution_text.get("1.0", "end-1c").strip()
        if not resolution:
            self.status_var.set("Escribe el texto de Resolución.")
            self._append_log("Falta ingresar el texto de Resolución.")
            self.resolution_text.focus_set()
            return

        self._clear_results_table()
        self._append_log(f"{incident}: iniciado.")
        self._queue_incident(incident, resolver_group, resolver_user, resolution)

    def _start_excel_batch(self) -> None:
        if self.busy or self.closing:
            return
        if not self.loaded_incidents:
            self.status_var.set("Carga un Excel con incidencias válidas.")
            self._append_log("No hay incidencias cargadas para procesar.")
            return

        resolver_group = self.resolver_group_var.get()
        resolver_user = self.resolver_user_var.get()
        if resolver_group not in RESOLVER_USERS:
            self.status_var.set("Selecciona un Grupo Resolutor.")
            self._append_log("Falta seleccionar el Grupo Resolutor para procesar el Excel.")
            self.resolver_group_combo.focus_set()
            return
        if resolver_user not in RESOLVER_USERS[resolver_group]:
            self.status_var.set("Selecciona un usuario válido para el grupo elegido.")
            self._append_log("Falta seleccionar un Usuario Resolutor válido para procesar el Excel.")
            self.resolver_user_combo.focus_set()
            return
        resolution = self.resolution_text.get("1.0", "end-1c").strip()
        if not resolution:
            self.status_var.set("Escribe el texto de Resolución.")
            self._append_log("Falta ingresar el texto de Resolución para procesar el Excel.")
            self.resolution_text.focus_set()
            return

        if not messagebox.askyesno(
            "Confirmar cierre masivo",
            "Se ejecutará el cierre masivo con estos datos:\n\n"
            f"Incidencias: {len(self.loaded_incidents)}\n"
            f"Grupo: {resolver_group}\n"
            f"Usuario: {resolver_user}\n\n"
            "¿Continuar?",
        ):
            self.status_var.set("Cierre masivo cancelado por el usuario.")
            self._append_log("Cierre masivo cancelado.")
            return

        self.batch_running = True
        self.batch_index = 0
        self.batch_results = []
        self._clear_results_table()
        self.batch_progress.configure(maximum=len(self.loaded_incidents), value=0)
        self.batch_progress_var.set(f"Procesando 0/{len(self.loaded_incidents)}")
        self._append_log(f"Cierre masivo iniciado: {len(self.loaded_incidents)} incidencias.")
        self._submit_next_batch_incident(resolver_group, resolver_user, resolution)

    def _submit_next_batch_incident(
        self, resolver_group: str, resolver_user: str, resolution: str
    ) -> None:
        if self.batch_index >= len(self.loaded_incidents):
            self._finish_batch()
            return

        incident = self.loaded_incidents[self.batch_index]
        self.batch_progress_var.set(
            f"Procesando {self.batch_index + 1}/{len(self.loaded_incidents)}"
        )
        self._append_log(f"{incident}: iniciado.")
        self._queue_incident(incident, resolver_group, resolver_user, resolution)

    def _queue_incident(
        self,
        incident: str,
        resolver_group: str,
        resolver_user: str,
        resolution: str,
    ) -> None:
        self.busy = True
        self._set_busy_buttons()
        self.status_var.set(
            f"Automatización en proceso: {incident}. No manipular Remedy hasta finalizar."
        )
        self.commands.put(
            WorkerCommand(
                action="search",
                incident=incident,
                resolver_group=resolver_group,
                resolver_user=resolver_user,
                resolution=resolution,
            )
        )

    def _set_idle_buttons(self) -> None:
        state = "disabled" if self.closing else "normal"
        self.search_button.configure(state=state)
        self.load_excel_button.configure(state=state)
        self.process_excel_button.configure(
            state="normal" if state == "normal" and self.loaded_incidents else "disabled"
        )

    def _set_busy_buttons(self) -> None:
        self.search_button.configure(state="disabled")
        self.load_excel_button.configure(state="disabled")
        self.process_excel_button.configure(state="disabled")

    def _finish_current_operation(self) -> None:
        self.busy = False
        self._set_idle_buttons()

    def _clear_results_table(self) -> None:
        for item in self.results_table.get_children():
            self.results_table.delete(item)

    def _append_result_row(self, incident: str, status: BatchStatus, detail: str) -> None:
        row_number = len(self.results_table.get_children()) + 1
        self.results_table.insert(
            "",
            "end",
            values=(row_number, incident, status, self._shorten_detail(detail)),
        )

    @staticmethod
    def _shorten_detail(message: str) -> str:
        normalized = " ".join(message.split())
        replacements = (
            (
                "Se pulsó Guardar, pero no se detectó confirmación; verifica Remedy antes de reintentar.",
                "Guardada; confirmación visible no detectada.",
            ),
            ("Remedy confirmó el guardado.", "Guardada."),
        )
        for source, target in replacements:
            if source in normalized:
                return target
        if "Helix redirigió a autenticación" in normalized:
            return "Sesión expirada o redirigida a login."
        if "campo editable" in normalized and "ID" in normalized:
            return "No se abrió el campo de búsqueda de incidencia."
        if "Redes Chile" in normalized:
            return "No se encontró o abrió la pestaña Redes Chile."
        if "Guardar está deshabilitado" in normalized:
            return "No se pudo guardar: botón Guardar deshabilitado."
        if "diálogo/modal inesperado" in normalized:
            return normalized[:180]
        return normalized[:180] or "Sin detalle."

    def _handle_incident_finished(
        self, incident: str | None, status: BatchStatus, message: str
    ) -> None:
        if not self.batch_running:
            self._finish_current_operation()
            if incident:
                self._clear_results_table()
                self._append_result_row(incident, status, message)
            if status == "OK":
                self._append_log(f"{incident}: cierre OK.")
                messagebox.showinfo(
                    "Trabajo finalizado",
                    f"Cierre manual finalizado para {incident}.",
                )
            else:
                self._append_log(
                    f"{incident or 'Incidencia'}: ERROR - {self._shorten_detail(message)}"
                )
                messagebox.showerror(
                    "Error en cierre manual",
                    f"No se pudo completar {incident or 'la incidencia'}.\n\n{message}",
                )
            return

        finished_incident = incident or self.loaded_incidents[self.batch_index]
        self.batch_results.append((finished_incident, status, message))
        self._append_result_row(finished_incident, status, message)
        self.batch_index += 1
        self.batch_progress.configure(value=self.batch_index)
        if status == "ERROR":
            self.batch_running = False
            report_path = self._write_batch_report()
            self.status_var.set(
                f"Excel detenido en {finished_incident}. Revisa el error antes de continuar."
            )
            self.batch_progress_var.set(
                f"Detenido: {self.batch_index}/{len(self.loaded_incidents)}"
            )
            self._append_log(
                f"{finished_incident}: ERROR - {self._shorten_detail(message)}"
            )
            self._append_log("Cierre masivo detenido por seguridad.")
            self._append_log(f"Reporte parcial generado: {report_path}")
            self._finish_current_operation()
            messagebox.showerror(
                "Cierre masivo detenido",
                f"No se pudo completar {finished_incident}.\n\n{message}",
            )
            return

        self._append_log(f"{finished_incident}: cierre OK.")
        resolver_group = self.resolver_group_var.get()
        resolver_user = self.resolver_user_var.get()
        resolution = self.resolution_text.get("1.0", "end-1c").strip()
        self.busy = False
        self._submit_next_batch_incident(resolver_group, resolver_user, resolution)

    def _finish_batch(self) -> None:
        self.batch_running = False
        report_path = self._write_batch_report()
        ok_count = sum(1 for _incident, status, _message in self.batch_results if status == "OK")
        error_count = len(self.batch_results) - ok_count
        self.status_var.set(
            f"Excel finalizado: {ok_count} OK, {error_count} con error. Revisa el reporte."
        )
        self.batch_progress.configure(value=len(self.loaded_incidents))
        self.batch_progress_var.set(f"Finalizado: {ok_count} OK, {error_count} errores")
        self._append_log(f"Cierre masivo finalizado: {ok_count} OK, {error_count} errores.")
        self._append_log(f"Reporte generado: {report_path}")
        self._finish_current_operation()
        messagebox.showinfo(
            "Trabajos finalizados",
            f"Cierre masivo finalizado.\nOK: {ok_count}\nErrores: {error_count}",
        )

    def _write_batch_report(self) -> Path:
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        now = _report_now()
        timestamp = now.strftime("%Y%m%d_%H%M%S")
        report_path = LOG_DIR / f"resultado_cierre_excel_{timestamp}.txt"
        with report_path.open("w", encoding="utf-8") as report:
            report.write(f"Cierres automatizados {now.strftime(REPORT_TITLE_FORMAT)}\n\n")
            report.write(f"Grupo Resolutor: {self.resolver_group_var.get()}\n")
            report.write(f"Usuario Resolutor: {self.resolver_user_var.get()}\n")
            report.write(
                "Resolución: "
                + " ".join(self.resolution_text.get("1.0", "end-1c").strip().split())
                + "\n\n"
            )
            ok_count = sum(1 for _incident, status, _message in self.batch_results if status == "OK")
            error_count = len(self.batch_results) - ok_count
            report.write("Resumen:\n")
            report.write(f"Total: {len(self.batch_results)}\n")
            report.write(f"OK: {ok_count}\n")
            report.write(f"Errores: {error_count}\n\n")
            report.write("Detalle:\n")
            report.write("numero de incidencia      Estado    Detalle\n")
            for incident, status, message in self.batch_results:
                detail = self._shorten_detail(message)
                report.write(f"{incident:<25}{status:<10}{detail}\n")
        return report_path

    def _poll_events(self) -> None:
        while True:
            try:
                event = self.events.get_nowait()
            except queue.Empty:
                break
            self._handle_event(event)
        if not self.closing:
            self.root.after(100, self._poll_events)

    def _handle_event(self, event: WorkerEvent) -> None:
        if event.kind == "log":
            return
        if event.kind == "ready":
            self.status_var.set("Lista. Edge se abrirá al enviar una incidencia.")
            self._append_log("Aplicación lista.")
            self._finish_current_operation()
        elif event.kind == "warning":
            self.status_var.set(event.message)
            self._append_log(f"Advertencia: {event.message}")
            self._finish_current_operation()
        elif event.kind == "done":
            self.status_var.set(f"Incidencia verificada: {event.incident}")
            self._handle_incident_finished(event.incident, "OK", event.message)
        elif event.kind == "searched":
            self.status_var.set(
                f"{event.incident}: Guardar enviado. Revisa la confirmación en Remedy."
            )
            self._handle_incident_finished(event.incident, "OK", event.message)
        elif event.kind == "error":
            self.status_var.set(f"No se pudo completar: {event.incident}.")
            self._handle_incident_finished(event.incident, "ERROR", event.message)
        elif event.kind == "closed":
            self.root.destroy()

    def _append_log(self, message: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", message + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _on_close(self) -> None:
        if self.closing:
            return
        self.closing = True
        self.busy = True
        self.batch_running = False
        self._set_busy_buttons()
        self.status_var.set("Cerrando Edge…")
        self.commands.put(WorkerCommand(action="stop"))
        self.root.after(100, self._poll_close)

    def _poll_close(self) -> None:
        while True:
            try:
                event = self.events.get_nowait()
            except queue.Empty:
                break
            self._handle_event(event)
            if event.kind == "closed":
                return
        self.root.after(100, self._poll_close)


def run_app() -> None:
    root = tk.Tk()
    HelixApp(root)
    root.mainloop()
