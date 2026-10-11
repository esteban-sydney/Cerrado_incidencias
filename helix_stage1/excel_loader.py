from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


INCIDENT_PATTERN = re.compile(r"INC\d{12}", re.IGNORECASE)
INCIDENT_HEADER_NAMES = {"incidencia", "incidencias"}
RESOLUTION_HEADER_NAMES = {"resolucion", "resoluciones"}


@dataclass(frozen=True)
class IncidentRecord:
    incident: str
    resolution: str


@dataclass(frozen=True)
class LoadResult:
    records: list[IncidentRecord]
    duplicates: list[str]
    invalid_values: list[str]
    missing_resolution: list[str]
    report_rows: list[tuple[str, str, str]]


def load_incident_file(path: Path) -> LoadResult:
    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        rows = _read_spreadsheet_rows_with_openpyxl(path)
    elif suffix == ".ods":
        rows = _read_spreadsheet_rows_with_ods(path)
    else:
        raise RuntimeError("Formato no soportado. Usa .xlsx, .xlsm o .ods.")

    return _extract_incident_records(rows)


def _read_spreadsheet_rows_with_openpyxl(path: Path) -> list[tuple[object, ...]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Falta instalar openpyxl. Ejecuta: pip install -r requirements.txt"
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def _read_spreadsheet_rows_with_ods(path: Path) -> list[tuple[object, ...]]:
    try:
        from pyexcel_ods3 import get_data
    except ImportError as error:
        raise RuntimeError(
            "Falta instalar pyexcel-ods3. Ejecuta: pip install -r requirements.txt"
        ) from error

    sheets = get_data(str(path))
    if not sheets:
        raise RuntimeError("El archivo .ods no contiene hojas.")
    first_sheet_rows = next(iter(sheets.values()))
    return [tuple(row) for row in first_sheet_rows]


def _extract_incident_records(rows: list[tuple[object, ...]]) -> LoadResult:
    header_index, incident_column, resolution_column = _find_headers(rows)

    records: list[IncidentRecord] = []
    seen: set[str] = set()
    duplicates: list[str] = []
    invalid_values: list[str] = []
    missing_resolution: list[str] = []
    report_rows: list[tuple[str, str, str]] = []

    for row in rows[header_index + 1 :]:
        raw_incident = _cell(row, incident_column)
        if raw_incident is None:
            continue

        incident = str(raw_incident).strip().upper()
        if not incident:
            continue

        raw_resolution = _cell(row, resolution_column)
        resolution = str(raw_resolution or "").strip()

        if INCIDENT_PATTERN.fullmatch(incident) is None:
            invalid_values.append(incident)
            report_rows.append((incident, "NOK", "Incidencia inválida"))
            continue
        if incident in seen:
            duplicates.append(incident)
            report_rows.append((incident, "NOK", "Incidencia duplicada"))
            continue
        if not resolution:
            missing_resolution.append(incident)
            report_rows.append((incident, "NOK", "Resolución vacía"))
            continue

        seen.add(incident)
        records.append(IncidentRecord(incident=incident, resolution=resolution))
        report_rows.append((incident, "OK", "Lista para procesar"))

    if not records:
        raise RuntimeError(
            "El archivo no contiene incidencias válidas con resolución. "
            "Debe incluir columnas 'Incidencias' y 'Resolución'."
        )
    return LoadResult(
        records=records,
        duplicates=duplicates,
        invalid_values=invalid_values,
        missing_resolution=missing_resolution,
        report_rows=report_rows,
    )


def _find_headers(rows: list[tuple[object, ...]]) -> tuple[int, int, int]:
    for row_index, row in enumerate(rows[:10]):
        incident_column: int | None = None
        resolution_column: int | None = None
        for column_index, value in enumerate(row):
            normalized = _normalize_header(value)
            if normalized in INCIDENT_HEADER_NAMES:
                incident_column = column_index
            elif normalized in RESOLUTION_HEADER_NAMES:
                resolution_column = column_index

        if incident_column is not None and resolution_column is not None:
            return row_index, incident_column, resolution_column

    raise RuntimeError("No encontré las columnas 'Incidencias' y 'Resolución'.")


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().casefold()
    replacements = str.maketrans("áéíóúüñ", "aeiouun")
    text = text.translate(replacements)
    return re.sub(r"[^a-z0-9]+", "", text)


def _cell(row: tuple[object, ...], index: int) -> object | None:
    if index >= len(row):
        return None
    return row[index]
